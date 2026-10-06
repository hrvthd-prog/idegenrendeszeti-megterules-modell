# -*- coding: utf-8 -*-
"""Idegenrendeszeti in-house ugyintezes megterulesi modell -> ures xlsx sablon."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

HUF = '#,##0" Ft"'
PCT = '0.0%'
NUM = '#,##0.0'
INT = '#,##0'

INPUT = PatternFill('solid', fgColor='FFF2CC')
CALC = PatternFill('solid', fgColor='EDEDED')
HEAD = PatternFill('solid', fgColor='1F4E79')
GOOD = PatternFill('solid', fgColor='E2EFDA')
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
GREY = Font(italic=True, size=9, color='808080')


def head(ws, row, labels, width=None, h=42):
    for i, t in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF', size=10)
        c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.row_dimensions[row].height = h
    if width:
        for i, w in enumerate(width, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, row, text, size=14):
    ws.cell(row=row, column=1, value=text).font = Font(bold=True, size=size, color='1F4E79')


wb = openpyxl.Workbook()

# ===================================================== 1. UTMUTATO
ws = wb.active
ws.title = 'Útmutató'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 112
ws.sheet_view.showGridLines = False
rows = [
    ('t', 'Idegenrendészeti ügyintézés házon belül – megtérülési kimutatás'),
    ('', ''),
    ('n', 'Mit mutat ez a munkafüzet?'),
    ('b', 'Azt a pénzbeli megtakarítást, ami abból keletkezik, hogy az idegenrendészeti ügyintézést a vállalat '
          'házon belül végzi, és nem külső szolgáltatónak fizet szolgáltatási díjat. A megtakarítást a pozíció '
          'teljes munkáltatói költségével állítja szembe.'),
    ('', ''),
    ('n', 'Használat – 3 lépés'),
    ('b', '1.  „Adatok” lap: töltsd ki a SÁRGA cellákat soronként (ügytípus, érintett fő, hosszabbítási ciklus, '
          'külső szolgáltatói díj / ügy). A szürke cellák számítottak – ne írd át őket.'),
    ('b', '2.  „Összegzés” lap, 1. blokk: írd be a bruttó bért, a hónapok számát, a szocho kulcsot és az egyéb '
          'munkáltatói ráfordítást.'),
    ('b', '3.  Ennyi. A 2–4. blokk és a diagram automatikusan frissül; a prezentációba ezekből a számokból '
          'kell átvenni.'),
    ('', ''),
    ('n', 'Módszertani döntések – ezekre egy pénzügyes rá fog kérdezni'),
    ('b', '•  Évesítés:  a 2–4 évente hosszabbítandó ügyeket nem a tényleges évükben számoljuk el, hanem elosztjuk '
          'a ciklus hosszával (pl. 3 éves ciklus esetén az érintett fő egyharmada esik egy évre). Így nincs '
          'mesterségesen jó vagy gyenge év, és nem vádolható a kimutatás cseresznyézéssel.'),
    ('b', '•  Eljárási illeték:  az államnak fizetett illeték KI VAN VÉVE a megtakarításból, mert azt belső és '
          'külső ügyintézés esetén egyaránt fizetni kell. Csak tájékoztató oszlopként szerepel. Ez lenne a '
          'kimutatás legtámadhatóbb pontja – ezért van előre kizárva.'),
    ('b', '•  Költségoldal:  teljes munkáltatói költség (bruttó bér + szociális hozzájárulási adó + egyéb '
          'munkáltatói ráfordítás), nem csak a bruttó bér. Ez a korrekt összehasonlítás.'),
    ('b', '•  Csak kemény számok:  elkerült bírság, felszabaduló HR-idő, rövidebb átfutási idő NINCS beszámítva. '
          'Ezek valós előnyök, de nem készpénz jellegűek – ha a modell nélkülük is pozitív, erősebb az állítás. '
          'Érvként szóban elmondhatók, számként nem szerepelnek.'),
    ('', ''),
    ('n', 'Amit a bemutatás előtt érdemes alátámasztásként becsatolni'),
    ('b', '•  A korábbi külső szolgáltató árlistája vagy a kiállított számlák → a „Díj / ügy” oszlop alátámasztása.'),
    ('b', '•  Az ügyintézési napló vagy a kezelt ügyek listája → az „Érintett fő” és a „Tény ügyszám” oszlopok '
          'alátámasztása.'),
    ('b', '•  A bérköltség-adatok HR / bérszámfejtés általi visszaigazolása.'),
    ('', ''),
    ('n', 'Színkód'),
    ('b', 'SÁRGA = kitöltendő bemeneti adat        SZÜRKE = számított, ne írd át        ZÖLD = végeredmény'),
]
r = 1
for kind, text in rows:
    if kind == 't':
        title(ws, r, text, 15)
    elif kind == 'n':
        ws.cell(row=r, column=2, value=text).font = Font(bold=True, size=11, color='1F4E79')
    elif kind == 'b':
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = 14 * (1 + len(text) // 105)
    r += 1

# ===================================================== 2. ADATOK
d = wb.create_sheet('Adatok')
d.sheet_view.showGridLines = False
title(d, 1, 'Ügytípusok, darabszámok és külső szolgáltatói díjak')
d.cell(row=2, column=1, value='Minden sárga cella kitöltendő. Az ügytípus-megnevezések szabadon átírhatók, '
                              'a felsorolás csak kiindulópont. Üres sor nem zavar, a végösszeg kezeli.').font = GREY

HR = 4
head(d, HR, [
    'Ügytípus', 'Érintett fő', 'Hosszabbítási\nciklus (év)', 'Évesített\nügyszám\n(számított)',
    'Tény ügyszám\nelmúlt 12 hó', 'Külső szolgáltatói\ndíj / ügy',
    'Elkerült díj –\névesített', 'Elkerült díj –\ntény 12 hó',
    'Eljárási illeték / ügy\n(tájékoztató, nem\nmegtakarítás)', 'Megjegyzés / forrás',
], width=[48, 11, 13, 12, 13, 18, 18, 18, 20, 42])

nevek = [
    'Tartózkodási engedély – új kérelem (munkavállalási cél)',
    'Tartózkodási engedély – hosszabbítás (1 éves ciklus)',
    'Tartózkodási engedély – hosszabbítás (2 éves ciklus)',
    'Tartózkodási engedély – hosszabbítás (3 éves ciklus)',
    'Tartózkodási engedély – hosszabbítás (4 éves ciklus)',
    'Tartózkodási engedély – pótlás / csere / adatváltozás',
    'Szálláshely- és lakcímbejelentés, lakcímváltozás',
    'Adószám / TAJ-szám igénylése',
    'Útlevélcsere átvezetése',
    'Igazolások, egyéb hatósági ügyintézés',
    '', '', '', '', '',
]
first = HR + 1
for i, nev in enumerate(nevek):
    r = first + i
    d.cell(row=r, column=1, value=nev).fill = INPUT
    for col in (2, 3, 5, 6, 9, 10):
        d.cell(row=r, column=col).fill = INPUT
    c = d.cell(row=r, column=4, value='=IFERROR(B{0}/C{0},"")'.format(r))
    c.fill = CALC
    c.number_format = NUM
    c = d.cell(row=r, column=7, value='=IFERROR(ROUND(D{0}*F{0},0),"")'.format(r))
    c.fill = CALC
    c.number_format = HUF
    c = d.cell(row=r, column=8, value='=IFERROR(ROUND(E{0}*F{0},0),"")'.format(r))
    c.fill = CALC
    c.number_format = HUF
    for col in (2, 3, 5):
        d.cell(row=r, column=col).number_format = INT
    for col in (6, 9):
        d.cell(row=r, column=col).number_format = HUF
    for col in range(1, 11):
        d.cell(row=r, column=col).border = BOX

last = first + len(nevek) - 1
TR = last + 1
d.cell(row=TR, column=1, value='ÖSSZESEN').font = Font(bold=True)
sums = {
    2: '=SUM(B{0}:B{1})'.format(first, last),
    4: '=SUM(D{0}:D{1})'.format(first, last),
    5: '=SUM(E{0}:E{1})'.format(first, last),
    7: '=SUM(G{0}:G{1})'.format(first, last),
    8: '=SUM(H{0}:H{1})'.format(first, last),
    9: '=SUMPRODUCT(IFERROR(D{0}:D{1},0),IFERROR(I{0}:I{1},0))'.format(first, last),
}
for col, f in sums.items():
    c = d.cell(row=TR, column=col, value=f)
    c.font = Font(bold=True)
    c.fill = GOOD
    c.border = BOX
    c.number_format = NUM if col == 4 else (HUF if col in (7, 8, 9) else INT)
d.cell(row=TR + 2, column=1, value='Az illeték-összesen évesített ügyszámmal számolt tájékoztató adat – '
                                   'a megtakarításban NEM szerepel, mert külső ügyintézés esetén is fizetni kellene.').font = GREY
d.freeze_panes = 'A{0}'.format(first)

# ===================================================== 3. OSSZEGZES
s = wb.create_sheet('Összegzés')
s.sheet_view.showGridLines = False
s.column_dimensions['A'].width = 54
for col in 'BCD':
    s.column_dimensions[col].width = 20
s.column_dimensions['E'].width = 56

title(s, 1, 'Megtérülési kimutatás – idegenrendészeti ügyintézés házon belül')
s.cell(row=2, column=1, value='Minden összeg HUF. Csak a sárga cellákat kell kitölteni.').font = GREY

B0 = 4
s.cell(row=B0, column=1, value='1.  Költségoldal – teljes munkáltatói költség').font = Font(bold=True, size=12, color='1F4E79')
ber = [
    ('Bruttó alapbér / hó', None, HUF, 'A HR / bérszámfejtés által visszaigazolt összeg'),
    ('Figyelembe vett hónapok / év', 12, INT, 'Ha van 13. havi vagy bónusz, itt növeld (pl. 13)'),
    ('Szociális hozzájárulási adó (szocho)', 0.13, PCT, 'Az aktuális kulcs – ellenőrizd a bérszámfejtésnél'),
    ('Egyéb munkáltatói ráfordítás / év', None, HUF, 'Eszköz, telefon, képzés, utazás – ha nincs adat, 0'),
]
for i, (nev, val, fmt, megj) in enumerate(ber):
    r = B0 + 1 + i
    s.cell(row=r, column=1, value=nev)
    c = s.cell(row=r, column=2, value=val)
    c.number_format = fmt
    c.border = BOX
    c.fill = INPUT
    s.cell(row=r, column=5, value=megj).font = GREY
BRUTTO, HONAP, SZOCHO, EGYEB = B0 + 1, B0 + 2, B0 + 3, B0 + 4
TOT = B0 + 5
s.cell(row=TOT, column=1, value='Teljes munkáltatói költség / év').font = Font(bold=True)
c = s.cell(row=TOT, column=2, value='=ROUND(B{0}*B{1}*(1+B{2})+B{3},0)'.format(BRUTTO, HONAP, SZOCHO, EGYEB))
c.number_format = HUF
c.font = Font(bold=True)
c.fill = CALC
c.border = BOX
s.cell(row=TOT, column=5, value='= bruttó × hónapok × (1 + szocho) + egyéb ráfordítás').font = GREY

E0 = TOT + 2
s.cell(row=E0, column=1, value='2.  Eredmény – évesített nézet').font = Font(bold=True, size=12, color='1F4E79')
SAV = E0 + 1
s.cell(row=SAV, column=1, value='Elkerült külső szolgáltatási díj (évesített)')
c = s.cell(row=SAV, column=2, value="='Adatok'!G{0}".format(TR))
c.number_format = HUF
c.fill = CALC
c.border = BOX
COST = E0 + 2
s.cell(row=COST, column=1, value='Teljes munkáltatói költség (mínusz)')
c = s.cell(row=COST, column=2, value='=-B{0}'.format(TOT))
c.number_format = HUF
c.fill = CALC
c.border = BOX
NET = E0 + 3
s.cell(row=NET, column=1, value='NETTÓ MEGTAKARÍTÁS / ÉV').font = Font(bold=True, size=12)
c = s.cell(row=NET, column=2, value='=B{0}+B{1}'.format(SAV, COST))
c.number_format = HUF
c.font = Font(bold=True, size=12)
c.fill = GOOD
c.border = BOX
ROI = E0 + 4
s.cell(row=ROI, column=1, value='Megtérülés  (nettó megtakarítás / munkáltatói költség)')
c = s.cell(row=ROI, column=2, value='=IFERROR(B{0}/B{1},"")'.format(NET, TOT))
c.number_format = PCT
c.fill = CALC
c.border = BOX
s.cell(row=ROI, column=5, value='Minden 1 Ft bérköltségre ennyi arányú nettó megtakarítás jut').font = GREY
COV = E0 + 5
s.cell(row=COV, column=1, value='Fedezet  (elkerült díj / munkáltatói költség)')
c = s.cell(row=COV, column=2, value='=IFERROR(B{0}/B{1},"")'.format(SAV, TOT))
c.number_format = PCT
c.fill = CALC
c.border = BOX
s.cell(row=COV, column=5, value='100% felett a pozíció kigazdálkodja önmagát').font = GREY
BEP = E0 + 6
s.cell(row=BEP, column=1, value='Fedezeti ügyszám – ennyi ügy hozná nullára')
c = s.cell(row=BEP, column=2, value="=IFERROR(B{0}/('Adatok'!G{1}/'Adatok'!D{1}),\"\")".format(TOT, TR))
c.number_format = NUM
c.fill = CALC
c.border = BOX
s.cell(row=BEP, column=5, value='Átlagos ügydíjjal számolva – hasonlítsd a lenti tényleges ügyszámhoz').font = GREY
ACT = E0 + 7
s.cell(row=ACT, column=1, value='Tényleges évesített ügyszám')
c = s.cell(row=ACT, column=2, value="='Adatok'!D{0}".format(TR))
c.number_format = NUM
c.fill = CALC
c.border = BOX

T0 = ACT + 2
s.cell(row=T0, column=1, value='3.  Kontrollnézet – tényleges elmúlt 12 hónap').font = Font(bold=True, size=12, color='1F4E79')
s.cell(row=T0 + 1, column=1, value='Elkerült külső szolgáltatási díj (tény 12 hó)')
c = s.cell(row=T0 + 1, column=2, value="='Adatok'!H{0}".format(TR))
c.number_format = HUF
c.fill = CALC
c.border = BOX
s.cell(row=T0 + 2, column=1, value='Nettó megtakarítás (tény 12 hó)').font = Font(bold=True)
c = s.cell(row=T0 + 2, column=2, value='=B{0}-B{1}'.format(T0 + 1, TOT))
c.number_format = HUF
c.font = Font(bold=True)
c.fill = GOOD
c.border = BOX
s.cell(row=T0 + 2, column=5, value='Csak a megtörtént ügyek – a legkevésbé támadható szám. '
                                   'Ha ez és az évesített nézet is pozitív, nincs mit vitatni.').font = GREY

K0 = T0 + 4
s.cell(row=K0, column=1, value='4.  Hároméves kumulált hatás (évesített alapon)').font = Font(bold=True, size=12, color='1F4E79')
IDX = K0 + 1
s.cell(row=IDX, column=1, value='Éves bérindexálás')
c = s.cell(row=IDX, column=2)
c.number_format = PCT
c.fill = INPUT
c.border = BOX
FIDX = K0 + 2
s.cell(row=FIDX, column=1, value='Külső szolgáltatói díjak éves emelkedése')
c = s.cell(row=FIDX, column=2)
c.number_format = PCT
c.fill = INPUT
c.border = BOX
s.cell(row=FIDX, column=5, value='Ha a két indexet egyenlőre állítod, az a semleges feltételezés. '
                                 'Ha a bérindexet magasabbra, az a konzervatív (óvatos) változat.').font = GREY

KH = K0 + 4
head(s, KH, ['Év', 'Elkerült díj', 'Munkáltatói költség', 'Nettó megtakarítás'], h=28)
for i in range(3):
    r = KH + 1 + i
    s.cell(row=r, column=1, value='{0}. év'.format(i + 1)).font = Font(bold=True)
    c = s.cell(row=r, column=2, value='=ROUND($B${0}*(1+$B${1})^{2},0)'.format(SAV, FIDX, i))
    c.number_format = HUF
    c.fill = CALC
    c.border = BOX
    c = s.cell(row=r, column=3, value='=ROUND($B${0}*(1+$B${1})^{2},0)'.format(TOT, IDX, i))
    c.number_format = HUF
    c.fill = CALC
    c.border = BOX
    c = s.cell(row=r, column=4, value='=B{0}-C{0}'.format(r))
    c.number_format = HUF
    c.fill = CALC
    c.border = BOX
KT = KH + 4
s.cell(row=KT, column=1, value='3 év összesen').font = Font(bold=True)
for col in (2, 3, 4):
    L = get_column_letter(col)
    c = s.cell(row=KT, column=col, value='=SUM({0}{1}:{0}{2})'.format(L, KH + 1, KH + 3))
    c.number_format = HUF
    c.font = Font(bold=True)
    c.fill = GOOD
    c.border = BOX

ch = BarChart()
ch.type = 'col'
ch.title = 'Elkerült szolgáltatási díj vs. munkáltatói költség'
ch.y_axis.title = 'HUF'
ch.height, ch.width = 8.5, 18
ch.add_data(Reference(s, min_col=2, max_col=3, min_row=KH, max_row=KH + 3), titles_from_data=True)
ch.set_categories(Reference(s, min_col=1, min_row=KH + 1, max_row=KH + 3))
s.add_chart(ch, 'A{0}'.format(KT + 3))

out = 'Megterulesi-kimutatas-idegenrendeszet.xlsx'
wb.save(out)
print('OK ->', out)
print('Adatok osszesen sor =', TR, '| Osszegzes netto sor =', NET, '| 3 ev osszesen sor =', KT)
